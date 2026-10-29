"""Builds Nurse-Schedule.xlsx: a formula-only (no macros) nursing schedule and
coverage calendar for Microsoft 365 Excel (desktop or web).

    python3 excel/build_workbook.py [output.xlsx]

Input sheets: Setup, Shifts, Recurring, Changes. Output sheets: Calendar,
Agenda, Attention. Hidden helper sheets (_Entries, _Slots, _Req, _Lists) do
the work: every one-time shift and every recurring occurrence that falls in
the 6-week calendar window becomes one row in _Entries, and everything else
(calendar text, overlaps, coverage gaps, staffing checks) reads from there.
"""

import datetime as dt
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

OUT = sys.argv[1] if len(sys.argv) > 1 else "excel/Nurse-Schedule.xlsx"

# Capacities
N_SHIFTS = 300       # one-time shift rows
N_SERIES = 30        # recurring schedules
N_CHANGES = 200      # per-date / from-date changes
N_NURSES = 60
N_FAC = 15
N_REQ = 8            # staffing requirements
GRID = 42            # calendar window (6 weeks)
N_AGENDA = 600
N_ATT = 100
SLOTS_PER_DAY = 48   # 30-minute staffing checks
N_SLOTS = (GRID + 1) * SLOTS_PER_DAY
EPS = "0.0003"       # ~26 s, keeps time comparisons away from float edges

FIRST = 6            # first data row on input sheets (header on row 5)
S_LAST = FIRST + N_SHIFTS - 1
R_LAST = FIRST + N_SERIES - 1
C_LAST = FIRST + N_CHANGES - 1
NUR_LAST = FIRST + N_NURSES - 1
FAC_LAST = FIRST + N_FAC - 1
REQ_FIRST = 6
REQ_LAST = REQ_FIRST + N_REQ - 1

E_S0 = 2                          # _Entries rows mirroring Shifts
E_R0 = E_S0 + N_SHIFTS            # _Entries rows for recurring candidates
E_LAST = E_R0 + N_SERIES * GRID - 1

CANCEL = "Cancel this date"
CHANGE = "Change this date"
RELEASE = "Release this date (needs coverage)"
FROM = "Change from this date on"
ACTIONS = [CANCEL, CHANGE, RELEASE, FROM]

FONT = "Arial"
INK = "1C2430"
MUTED = "5B6676"
BLUE_IN = "0000FF"
HEAD_FILL = PatternFill("solid", fgColor="E8EEF8")
INPUT_FILL = PatternFill("solid", fgColor="FFFDF0")
CALC_FILL = PatternFill("solid", fgColor="F3F4F6")
THIN = Side(style="thin", color="D5DAE1")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def font(**kw):
    kw.setdefault("name", FONT)
    kw.setdefault("color", INK)
    return Font(**kw)


wb = Workbook()
ws_cal = wb.active
ws_cal.title = "Calendar"
ws_ag = wb.create_sheet("Agenda")
ws_att = wb.create_sheet("Attention")
ws_sh = wb.create_sheet("Shifts")
ws_rec = wb.create_sheet("Recurring")
ws_ch = wb.create_sheet("Changes")
ws_set = wb.create_sheet("Setup")
ws_help = wb.create_sheet("Instructions")
ws_e = wb.create_sheet("_Entries")
ws_sl = wb.create_sheet("_Slots")
ws_rq = wb.create_sheet("_Req")
ws_li = wb.create_sheet("_Lists")


def name(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def rng(sheet, col, a, b):
    return f"{sheet}!${col}${a}:${col}${b}"


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = font(size=16, bold=True, color="2563EB")
    ws["A2"] = sub
    ws["A2"].font = font(size=10, color=MUTED, italic=True)


def header(ws, row, labels, widths=None, start_col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row, start_col + i, lab)
        c.font = font(bold=True, size=10)
        c.fill = HEAD_FILL
        c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
        if widths:
            ws.column_dimensions[L(start_col + i)].width = widths[i]
    ws.row_dimensions[row].height = 30


def style_rows(ws, r0, r1, c0, c1, fill=None, fmt=None, color=None, wrap=False):
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            cell = ws.cell(r, c)
            cell.font = font(size=10, color=color or INK)
            cell.border = BOX
            if fill:
                cell.fill = fill
            if fmt:
                cell.number_format = fmt
            if wrap:
                cell.alignment = Alignment(wrap_text=True, vertical="top")


def add_list(ws, ref_range, formula, prompt=None, allow_blank=True, error=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=allow_blank, showErrorMessage=error)
    if prompt:
        dv.promptTitle, dv.prompt, dv.showInputMessage = "Tip", prompt, True
    ws.add_data_validation(dv)
    dv.add(ref_range)


# ---------------------------------------------------------------------------
# _Lists
# ---------------------------------------------------------------------------
li = ws_li
li["A1"] = "Pts (15-min offsets)"
for k in range(96):
    li.cell(2 + k, 1, k)
name("Pts", "_Lists!$A$2:$A$97")
li["B1"] = "WdNums →"
for i in range(7):
    li.cell(1, 3 + i, i + 1)
    li.cell(2, 3 + i, ["Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"][i])
name("WdNums", "_Lists!$C$1:$I$1")
name("WdNames", "_Lists!$C$2:$I$2")
li["K1"] = "Active nurses"
for k in range(N_NURSES):
    r = 2 + k
    li.cell(r, 11, f'=IFERROR(INDEX(Setup!$A${FIRST}:$A${NUR_LAST},MATCH({k + 1},Setup!$C${FIRST}:$C${NUR_LAST},0)),"")')
name("ActiveNurses", f"_Lists!$K$2:$K${N_NURSES + 1}")
li["M1"] = "Facility filter"
li["M2"] = "All"
for k in range(N_FAC):
    li.cell(3 + k, 13, f'=IF(Setup!$F${FIRST + k}="","",Setup!$F${FIRST + k})')
li["N1"] = "Nurse filter"
li["N2"] = "All"
for k in range(N_NURSES):
    li.cell(3 + k, 14, f'=IF(Setup!$A${FIRST + k}="","",Setup!$A${FIRST + k})')
for col, (head, vals) in zip("PQRST", [
    ("Types", ["Regular", "Coverage"]),
    ("Statuses", ["Confirmed", "Tentative", "Unfilled", "Released"]),
    ("Rec statuses", ["Confirmed", "Tentative", "Unfilled"]),
    ("Actions", ACTIONS),
    ("YN", ["Y"]),
]):
    li[f"{col}1"] = head
    for i, v in enumerate(vals):
        li[f"{col}{2 + i}"] = v
li.sheet_state = "hidden"

# ---------------------------------------------------------------------------
# Setup: nurses, facilities, requirements
# ---------------------------------------------------------------------------
st = ws_set
title(st, "Setup", "Nurses, facilities and optional staffing requirements. Yellow cells are inputs.")
st["A4"], st["F4"], st["I4"] = "Nurses", "Facilities", "Staffing requirements (optional)"
for c in ("A4", "F4", "I4"):
    st[c].font = font(bold=True, size=12)
header(st, 5, ["Nurse name", "Active? (Y/No)", "#"], [24, 12, 5])
header(st, 5, ["Facility name"], [26], start_col=6)
REQ_COLS = ["Facility", "Label", "Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat",
            "From", "To", "Min nurses", "Active? (Y/No)", "From slot", "To slot", "Overnight", "Length (slots)",
            "Effective from (optional)"]
header(st, 5, REQ_COLS, [22, 14, 5, 5, 5, 5, 5, 5, 5, 9, 9, 8, 9, 6, 6, 8, 8, 13], start_col=9)
style_rows(st, FIRST, NUR_LAST, 1, 2, fill=INPUT_FILL, color=BLUE_IN)
style_rows(st, FIRST, NUR_LAST, 3, 3, fill=CALC_FILL, color=MUTED)
for r in range(FIRST, NUR_LAST + 1):
    st.cell(r, 3, f'=IF(AND(A{r}<>"",B{r}<>"No"),COUNTIFS($A${FIRST}:A{r},"<>",$B${FIRST}:B{r},"<>No"),"")')
add_list(st, f"B{FIRST}:B{NUR_LAST}", '"Y,No"', "Choose No to hide a nurse from new assignments. Past assignments are kept.")
style_rows(st, FIRST, FAC_LAST, 6, 6, fill=INPUT_FILL, color=BLUE_IN)
# Requirements I..Y  (I fac, J label, K..Q days, R from, S to, T min, U active, V..Y computed)
style_rows(st, REQ_FIRST, REQ_LAST, 9, 21, fill=INPUT_FILL, color=BLUE_IN)
style_rows(st, REQ_FIRST, REQ_LAST, 22, 25, fill=CALC_FILL, color=MUTED)
style_rows(st, REQ_FIRST, REQ_LAST, 26, 26, fill=INPUT_FILL, color=BLUE_IN, fmt="mmm d, yyyy")
for r in range(REQ_FIRST, REQ_LAST + 1):
    st.cell(r, 18).number_format = "h:mm AM/PM"
    st.cell(r, 19).number_format = "h:mm AM/PM"
    st.cell(r, 22, f'=IF(R{r}="","",ROUND(R{r}*48,0))')
    st.cell(r, 23, f'=IF(S{r}="","",ROUND(S{r}*48,0))')
    st.cell(r, 24, f'=IF(OR(V{r}="",W{r}=""),"",W{r}<=V{r})')
    st.cell(r, 25, f'=IF(X{r}="","",W{r}-V{r}+IF(X{r},48,0))')
add_list(st, f"I{REQ_FIRST}:I{REQ_LAST}", f"=$F${FIRST}:$F${FAC_LAST}")
add_list(st, f"K{REQ_FIRST}:Q{REQ_LAST}", '"Y"')
add_list(st, f"U{REQ_FIRST}:U{REQ_LAST}", '"Y,No"')
st["I15"] = ("Requirements are checked every 30 minutes; use times on the hour or half hour. A 'To' time at or "
             "before 'From' means the window ends the next day. Only Confirmed assignments count; days met only "
             "with Tentative ones are flagged separately. Days without a requirement are never treated as gaps.")
st["I15"].font = font(size=9, color=MUTED, italic=True)
st.merge_cells("I15:Y17")
st["I15"].alignment = Alignment(wrap_text=True, vertical="top")
st["A3"] = "Don't sort or delete rows on the input sheets — clear a row's contents instead."
st["A3"].font = font(size=9, color="B91C1C", italic=True)
st.freeze_panes = "A6"

# ---------------------------------------------------------------------------
# Calendar inputs + defined names (used everywhere)
# ---------------------------------------------------------------------------
cal = ws_cal
name("MonthPick", "Calendar!$C$3")
name("FacFilter", "Calendar!$E$3")
name("NurseFilter", "Calendar!$G$3")
name("GridStart", "Calendar!$L$1")
name("GridEnd", "Calendar!$M$1")

# ---------------------------------------------------------------------------
# Shifts (one-time assignments)
# ---------------------------------------------------------------------------
sh = ws_sh
title(sh, "Shifts — one-time assignments",
      "One row per shift. For several dates, copy the row. Blank Type = Regular; blank Status = Confirmed.")
SH_COLS = ["ID", "Date", "Facility", "Nurse", "Start", "End", "Hours not specified (Y)", "Type",
           "Covering for", "Status", "Covers shift ID", "Notes", "Ends next day", "Hours", "Warnings",
           "Coverage still needed"]
header(sh, 5, SH_COLS, [6, 12, 20, 20, 10, 10, 11, 11, 18, 12, 16, 26, 9, 7, 34, 34])
style_rows(sh, FIRST, S_LAST, 1, 1, fill=CALC_FILL, color=MUTED)
style_rows(sh, FIRST, S_LAST, 2, 12, fill=INPUT_FILL, color=BLUE_IN)
style_rows(sh, FIRST, S_LAST, 13, 16, fill=CALC_FILL, color=MUTED, wrap=True)
for i in range(N_SHIFTS):
    r = FIRST + i
    e = E_S0 + i
    sh.cell(r, 1, f"S{i + 1}")
    sh.cell(r, 2).number_format = "ddd mmm d, yyyy"
    sh.cell(r, 5).number_format = "h:mm AM/PM"
    sh.cell(r, 6).number_format = "h:mm AM/PM"
    sh.cell(r, 13, f'=IF(OR(B{r}="",E{r}="",F{r}="",G{r}="Y"),"",IF(F{r}<=E{r},"Yes",""))')
    sh.cell(r, 14, f'=IF(_Entries!H{e},ROUND((_Entries!P{e}-_Entries!O{e})*24,2),"")')
    sh.cell(r, 14).number_format = "0.0"
    issues = (f'_xlfn.TEXTJOIN("; ",TRUE,'
              f'IF(AND(B{r}<>"",C{r}=""),"Choose a facility",""),'
              f'IF(AND(B{r}<>"",D{r}="",OR(J{r}="",J{r}="Confirmed",J{r}="Tentative")),"Confirmed/tentative shifts need a nurse",""),'
              f'IF(AND(B{r}<>"",G{r}<>"Y",OR(E{r}="",F{r}="")),"Enter start and end, or Y for hours not specified",""),'
              f'IF(AND(B{r}<>"",E{r}<>"",E{r}=F{r},G{r}<>"Y"),"Start and end are the same",""),'
              f'IF(AND(D{r}<>"",I{r}<>"",D{r}=I{r}),"A nurse cannot cover for themselves",""),'
              f'IF(AND(D{r}<>"",IFERROR(INDEX(Setup!$B${FIRST}:$B${NUR_LAST},MATCH(D{r},Setup!$A${FIRST}:$A${NUR_LAST},0))="No",FALSE)),"Nurse is inactive",""),'
              f'_Entries!T{e})')
    sh.cell(r, 15, f'=IF(B{r}="","",{issues})')
    sh.cell(r, 16, f'=IF(B{r}="","",_Entries!AE{e})')
add_list(sh, f"C{FIRST}:C{S_LAST}", f"=Setup!$F${FIRST}:$F${FAC_LAST}")
add_list(sh, f"D{FIRST}:D{S_LAST}", f"=_Lists!$K$2:$K${N_NURSES + 1}", error=False)
add_list(sh, f"I{FIRST}:I{S_LAST}", f"=Setup!$A${FIRST}:$A${NUR_LAST}")
add_list(sh, f"G{FIRST}:G{S_LAST}", '"Y"')
add_list(sh, f"H{FIRST}:H{S_LAST}", '"Regular,Coverage"')
add_list(sh, f"J{FIRST}:J{S_LAST}", '"Confirmed,Tentative,Unfilled,Released"',
         "Released = the assigned nurse is out and the shift needs coverage. Add coverage rows with this row's ID in 'Covers shift ID'.")
sh["K5"].comment = Comment("To cover (part of) another shift, put that shift's ID here: S12 for a one-time shift, "
                           "or the occurrence ID shown on the Changes sheet / Agenda for a recurring date.", "Nurse Schedule")
sh.freeze_panes = "C6"
sh["A3"] = "Don't sort or delete rows — clear a row's contents instead (IDs must stay put)."
sh["A3"].font = font(size=9, color="B91C1C", italic=True)

# ---------------------------------------------------------------------------
# Recurring schedules
# ---------------------------------------------------------------------------
rc = ws_rec
title(rc, "Recurring schedules",
      "One row per repeating assignment. Occurrences are calculated — no need to enter each date.")
RC_COLS = ["ID", "Facility", "Nurse", "Start", "End", "Hours not specified (Y)", "Type", "Status",
           "First date", "Every N weeks", "Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "End date (optional)",
           "Stop after N times (optional)", "Covering for", "Notes",
           "Interval", "Days/wk", "Week start", "Skipped in 1st week", "Summary"]
header(rc, 5, RC_COLS, [5, 18, 18, 10, 10, 10, 10, 11, 12, 8, 5, 5, 5, 5, 5, 5, 5, 12, 10, 16, 20, 7, 7, 11, 8, 70])
style_rows(rc, FIRST, R_LAST, 1, 1, fill=CALC_FILL, color=MUTED)
style_rows(rc, FIRST, R_LAST, 2, 21, fill=INPUT_FILL, color=BLUE_IN)
style_rows(rc, FIRST, R_LAST, 22, 26, fill=CALC_FILL, color=MUTED)
for i in range(N_SERIES):
    r = FIRST + i
    rc.cell(r, 1, f"R{i + 1}")
    for c in (4, 5):
        rc.cell(r, c).number_format = "h:mm AM/PM"
    for c in (9, 18):
        rc.cell(r, c).number_format = "ddd mmm d, yyyy"
    rc.cell(r, 22, f"=MAX(1,N(J{r}))")
    rc.cell(r, 23, f'=COUNTIF(K{r}:Q{r},"Y")')
    rc.cell(r, 24, f'=IF(I{r}="","",I{r}-WEEKDAY(I{r})+1)')
    rc.cell(r, 24).number_format = "mmm d, yyyy"
    rc.cell(r, 25, f'=IF(I{r}="","",SUMPRODUCT((K{r}:Q{r}="Y")*(WdNums<WEEKDAY(I{r}))))')
    hours = (f'IF(F{r}="Y",", hours not specified",IF(OR(D{r}="",E{r}=""),"",", "&TEXT(D{r},"h:mm AM/PM")&"–"'
             f'&TEXT(E{r},"h:mm AM/PM")&IF(E{r}<=D{r}," (next day)","")))')
    days = f'_xlfn.TEXTJOIN(", ",TRUE,IF(K{r}:Q{r}="Y",WdNames,""))'
    every = f'IF(V{r}=1,"Every ",IF(V{r}=2,"Every other week on ","Every "&V{r}&" weeks on "))'
    ends = (f'IF(R{r}<>"",", through "&TEXT(R{r},"mmmm d, yyyy"),'
            f'IF(S{r}<>"",", for "&S{r}&" occurrences",", with no end date"))')
    summary = (f'=IF(OR(B{r}="",I{r}=""),"",IF(W{r}=0,"Choose at least one weekday",'
               f'{every}&{days}&{hours}&", starting "&TEXT(I{r},"mmmm d, yyyy")&{ends}&"."))')
    rc[f"Z{r}"] = ArrayFormula(f"Z{r}", summary)
add_list(rc, f"B{FIRST}:B{R_LAST}", f"=Setup!$F${FIRST}:$F${FAC_LAST}")
add_list(rc, f"C{FIRST}:C{R_LAST}", f"=_Lists!$K$2:$K${N_NURSES + 1}", error=False)
add_list(rc, f"T{FIRST}:T{R_LAST}", f"=Setup!$A${FIRST}:$A${NUR_LAST}")
add_list(rc, f"F{FIRST}:F{R_LAST}", '"Y"')
add_list(rc, f"G{FIRST}:G{R_LAST}", '"Regular,Coverage"')
add_list(rc, f"H{FIRST}:H{R_LAST}", '"Confirmed,Tentative,Unfilled"')
add_list(rc, f"K{FIRST}:Q{R_LAST}", '"Y"')
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="52", allow_blank=True)
rc.add_data_validation(dv)
dv.add(f"J{FIRST}:J{R_LAST}")
rc.freeze_panes = "D6"
rc["A3"] = ("To stop a schedule, set its End date. To change one date or everything from a date onward, "
            "use the Changes sheet so earlier dates stay as they were.")
rc["A3"].font = font(size=9, color="B91C1C", italic=True)

# ---------------------------------------------------------------------------
# Changes (exceptions to recurring schedules)
# ---------------------------------------------------------------------------
cg = ws_ch
title(cg, "Changes to recurring schedules",
      "Cancel or change one date, release a date for coverage, or change everything from a date onward. "
      "Blank 'New' fields keep the schedule's value.")
CH_COLS = ["Schedule ID", "Date", "Action", "New nurse", "New start", "New end", "New status", "New type", "Notes",
           "Occurrence ID", "Check", "Single key", "Summary"]
header(cg, 5, CH_COLS, [10, 14, 30, 18, 10, 10, 12, 11, 22, 18, 26, 14, 50])
style_rows(cg, FIRST, C_LAST, 1, 9, fill=INPUT_FILL, color=BLUE_IN)
style_rows(cg, FIRST, C_LAST, 10, 13, fill=CALC_FILL, color=MUTED, wrap=True)
REC_ID = f"Recurring!$A${FIRST}:$A${R_LAST}"


def occurs_expr(sidx, d):
    """Boolean: does recurring row index `sidx` (1-based) occur on date `d`?"""
    def col(c):
        return f"INDEX(Recurring!${c}${FIRST}:${c}${R_LAST},{sidx})"
    n = (f"(INT(({d}-{col('X')})/7)/{col('V')}*{col('W')}"
         f"+SUMPRODUCT((INDEX(Recurring!$K${FIRST}:$Q${R_LAST},{sidx},0)=\"Y\")*(WdNums<=WEEKDAY({d})))-{col('Y')})")
    return (f'IFERROR(AND({col("B")}<>"",{col("I")}<>"",{d}>={col("I")},OR({col("R")}="",{d}<={col("R")}),'
            f'MOD(INT(({d}-{col("X")})/7),{col("V")})=0,'
            f'INDEX(Recurring!$K${FIRST}:$Q${R_LAST},{sidx},WEEKDAY({d}))="Y",'
            f'OR({col("S")}="",{n}<={col("S")})),FALSE)')


for i in range(N_CHANGES):
    r = FIRST + i
    cg.cell(r, 2).number_format = "ddd mmm d, yyyy"
    cg.cell(r, 5).number_format = "h:mm AM/PM"
    cg.cell(r, 6).number_format = "h:mm AM/PM"
    sidx = f"MATCH(A{r},{REC_ID},0)"
    cg.cell(r, 10, f'=IF(OR(A{r}="",B{r}=""),"",A{r}&" "&TEXT(B{r},"yyyy-mm-dd"))')
    cg.cell(r, 11, f'=IF(OR(A{r}="",B{r}="",C{r}=""),"",IF(ISNA({sidx}),"Unknown schedule ID",'
                   f'IF(C{r}="{FROM}","OK — applies from this date",IF({occurs_expr(sidx, f"B{r}")},"OK",'
                   f'"Not a scheduled date for "&A{r}))))')
    cg.cell(r, 12, f'=IF(OR(A{r}="",B{r}="",C{r}="",C{r}="{FROM}"),"",A{r}&"|"&TEXT(B{r},"0"))')
    cg.cell(r, 13, f'=IF(K{r}="","",C{r}&IF(D{r}<>""," · nurse: "&D{r},"")'
                   f'&IF(E{r}<>""," · start "&TEXT(E{r},"h:mm AM/PM"),"")&IF(F{r}<>""," · end "&TEXT(F{r},"h:mm AM/PM"),"")'
                   f'&IF(G{r}<>""," · "&G{r},"")&IF(H{r}<>""," · "&H{r},""))')
add_list(cg, f"A{FIRST}:A{C_LAST}", f"={REC_ID}")
add_list(cg, f"C{FIRST}:C{C_LAST}", f"=_Lists!$S$2:$S${1 + len(ACTIONS)}")
add_list(cg, f"D{FIRST}:D{C_LAST}", f"=_Lists!$K$2:$K${N_NURSES + 1}", error=False)
add_list(cg, f"G{FIRST}:G{C_LAST}", '"Confirmed,Tentative,Unfilled"')
add_list(cg, f"H{FIRST}:H{C_LAST}", '"Regular,Coverage"')
cg.column_dimensions["L"].hidden = True
cg.freeze_panes = "C6"

# ---------------------------------------------------------------------------
# _Entries: one row per one-time shift and per recurring occurrence in window
# ---------------------------------------------------------------------------
en = ws_e
E_COLS = {
    "A": "Src", "B": "ID", "C": "Date", "D": "Facility", "E": "Nurse", "F": "Start", "G": "End",
    "H": "HoursSpec", "I": "Type", "J": "Status", "K": "CoveringFor", "L": "Covers", "M": "Notes",
    "N": "Active", "O": "StartDT", "P": "EndDT", "Q": "Working", "R": "CountConf", "S": "CountAny",
    "T": "Overlap", "U": "Open", "V": "UncovConfH", "W": "UncovAnyH", "X": "FirstGap", "Y": "Label",
    "Z": "InWindow", "AA": "FacOK", "AB": "Key", "AC": "Rank", "AD": "NeedsCover", "AE": "CoverText",
    "AF": "AttRank",
    # recurring helpers
    "AG": "s", "AH": "d", "AI": "SeriesID", "AJ": "Occurs", "AK": "SingleKey", "AL": "ChRow", "AM": "ChAct",
    "AN": "NurseF", "AO": "StartF", "AP": "EndF", "AQ": "StatusF", "AR": "TypeF", "AS": "StartEff",
    "AT": "EndEff", "AU": "NurseOK",
}
for c, h in E_COLS.items():
    en[f"{c}1"] = h
    en[f"{c}1"].font = font(bold=True, size=9)
R = f"${E_S0}:$"  # noqa  (kept for readability below)


def E(col):
    return f"_Entries!${col}${E_S0}:${col}${E_LAST}"


def EL(col):
    return f"${col}${E_S0}:${col}${E_LAST}"


def common(r):
    """Formulas shared by both sections (columns O onward)."""
    f = {}
    f["O"] = f'=IF(NOT(N{r}),"",IF(H{r},C{r}+F{r},C{r}))'
    f["P"] = f'=IF(NOT(N{r}),"",IF(H{r},C{r}+G{r}+IF(G{r}<=F{r},1,0),C{r}))'
    f["Q"] = f'=AND(N{r},E{r}<>"",OR(J{r}="Confirmed",J{r}="Tentative"))'
    f["R"] = f'=AND(Q{r},H{r},J{r}="Confirmed")'
    f["S"] = f'=AND(Q{r},H{r})'
    f["T"] = (f'=IF(NOT(Q{r}),"",IF(H{r},'
              f'IF(COUNTIFS({EL("E")},E{r},{EL("S")},TRUE,{EL("O")},"<"&P{r},{EL("P")},">"&O{r})>1,"Overlapping assignment for "&E{r},'
              f'IF(COUNTIFS({EL("E")},E{r},{EL("Q")},TRUE,{EL("H")},FALSE,{EL("C")},C{r})>0,"Possible overlap (hours not specified)","")),'
              f'IF(COUNTIFS({EL("E")},E{r},{EL("Q")},TRUE,{EL("C")},C{r})>1,"Possible overlap (hours not specified)","")))')
    f["U"] = f'=AND(N{r},OR(J{r}="Unfilled",J{r}="Released"))'
    pts = f"Pts/96+{EPS}"
    inshift = f"(Pts<(P{r}-O{r})*96-0.01)"
    cov_conf = (f'COUNTIFS({EL("L")},B{r},{EL("Q")},TRUE,{EL("J")},"Confirmed",{EL("H")},TRUE,'
                f'{EL("O")},"<="&(O{r}+{pts}),{EL("P")},">"&(O{r}+{pts}))')
    cov_any = (f'COUNTIFS({EL("L")},B{r},{EL("Q")},TRUE,{EL("H")},TRUE,'
               f'{EL("O")},"<="&(O{r}+{pts}),{EL("P")},">"&(O{r}+{pts}))')
    whole_conf = f'COUNTIFS({EL("L")},B{r},{EL("Q")},TRUE,{EL("J")},"Confirmed",{EL("H")},FALSE)'
    whole_any = f'COUNTIFS({EL("L")},B{r},{EL("Q")},TRUE,{EL("H")},FALSE)'
    f["V"] = ArrayFormula(f"V{r}", f'=IF(AND(U{r},H{r}),IF({whole_conf}>0,0,SUMPRODUCT({inshift}*({cov_conf}=0))/4),"")')
    f["W"] = ArrayFormula(f"W{r}", f'=IF(AND(U{r},H{r}),IF({whole_any}>0,0,SUMPRODUCT({inshift}*({cov_any}=0))/4),"")')
    f["X"] = ArrayFormula(f"X{r}", f'=IF(AND(U{r},H{r}),IF(N(V{r})>0,MIN(IF({inshift}*({cov_conf}=0),O{r}+Pts/96)),""),"")')
    t = lambda x: f'SUBSTITUTE(SUBSTITUTE(TEXT({x},"h:mm AM/PM")," AM","a")," PM","p")'  # noqa: E731
    tag = (f'IF(J{r}="Released","RELEASED",IF(J{r}="Unfilled","UNFILLED",IF(I{r}="Coverage","Coverage","Regular")))'
           f'&IF(J{r}="Tentative"," · Tentative","")')
    f["Y"] = (f'=IF(NOT(N{r}),"",IF(H{r},{t(f"F{r}")}&"–"&{t(f"G{r}")}&IF(G{r}<=F{r}," (+1)",""),"Hours TBD")'
              f'&"  "&IF(E{r}="","—",E{r})&" · "&D{r}&" · "&{tag}&IF(K{r}<>""," · for "&K{r},"")'
              f'&IF(AD{r}," · NEEDS COVERAGE","")&IF(T{r}<>""," · ⚠ overlap",""))')
    f["Z"] = f"=AND(N{r},C{r}>=GridStart,C{r}<=GridEnd)"
    f["AA"] = f'=OR(FacFilter="All",FacFilter="",D{r}=FacFilter)'
    f["AU"] = f'=OR(NurseFilter="All",NurseFilter="",E{r}=NurseFilter,K{r}=NurseFilter)'
    f["AB"] = f'=IF(NOT(N{r}),"",IF(H{r},O{r},C{r}+0.999)+ROW()/10000000)'
    f["AC"] = f'=IF(AND(Z{r},AA{r},AU{r}),COUNTIFS({EL("Z")},TRUE,{EL("AA")},TRUE,{EL("AU")},TRUE,{EL("AB")},"<"&AB{r})+1,"")'
    f["AD"] = f'=AND(U{r},IF(H{r},N(V{r})>0,COUNTIFS({EL("L")},B{r},{EL("Q")},TRUE,{EL("J")},"Confirmed")=0))'
    f["AE"] = (f'=IF(NOT(U{r}),"",IF(NOT(AD{r}),"Covered",IF(H{r},"Still unfilled: "&TEXT(V{r},"0.##")&" h"'
               f'&IF(X{r}<>""," (first gap "&{t(f"X{r}")}&")","")'
               f'&IF(N(V{r})-N(W{r})>0,"; "&TEXT(V{r}-W{r},"0.##")&" h covered only tentatively",""),'
               f'"Needs a nurse (hours not specified)")))')
    f["AF"] = f'=IF(AND(AD{r},Z{r},AA{r}),COUNTIFS({EL("AD")},TRUE,{EL("Z")},TRUE,{EL("AA")},TRUE,{EL("AB")},"<"&AB{r})+1,"")'
    return f


for i in range(N_SHIFTS):
    r = E_S0 + i
    s = FIRST + i
    f = {
        "A": "S",
        "B": f'=IF(Shifts!B{s}="","",Shifts!A{s})',
        "C": f'=IF(Shifts!B{s}="","",Shifts!B{s})',
        "D": f'=Shifts!C{s}&""',
        "E": f'=IF(J{r}="Unfilled","",Shifts!D{s}&"")',
        "F": f'=IF(H{r},Shifts!E{s},"")',
        "G": f'=IF(H{r},Shifts!F{s},"")',
        "H": f'=AND(Shifts!G{s}<>"Y",Shifts!E{s}<>"",Shifts!F{s}<>"",Shifts!E{s}<>Shifts!F{s})',
        "I": f'=IF(Shifts!H{s}="","Regular",Shifts!H{s})',
        "J": f'=IF(Shifts!J{s}="","Confirmed",Shifts!J{s})',
        "K": (f'=IF(Shifts!I{s}<>"",Shifts!I{s},IF(Shifts!K{s}<>"",IFERROR(INDEX({EL("E")},'
              f'MATCH(Shifts!K{s},{EL("B")},0))&"",""),""))'),
        "L": f'=Shifts!K{s}&""',
        "M": f'=Shifts!L{s}&""',
        "N": f'=AND(Shifts!B{s}<>"",Shifts!C{s}<>"")',
    }
    f.update(common(r))
    for c, v in f.items():
        en[f"{c}{r}"] = v


def rcol(c, idx):
    return f"INDEX(Recurring!${c}${FIRST}:${c}${R_LAST},{idx})"


def chcol(c, idx):
    return f"INDEX(Changes!${c}${FIRST}:${c}${C_LAST},MAX(1,{idx}))"


def from_override(r, field_col, base):
    ld = (f'_xlfn.MAXIFS(Changes!$B${FIRST}:$B${C_LAST},Changes!$A${FIRST}:$A${C_LAST},AI{r},'
          f'Changes!$C${FIRST}:$C${C_LAST},"{FROM}",Changes!$B${FIRST}:$B${C_LAST},"<="&AH{r},'
          f'Changes!${field_col}${FIRST}:${field_col}${C_LAST},"<>")')
    pick = (f'INDEX(Changes!${field_col}${FIRST}:${field_col}${C_LAST},MATCH(1,(Changes!$A${FIRST}:$A${C_LAST}=AI{r})'
            f'*(Changes!$B${FIRST}:$B${C_LAST}={ld})*(Changes!$C${FIRST}:$C${C_LAST}="{FROM}")'
            f'*(Changes!${field_col}${FIRST}:${field_col}${C_LAST}<>""),0))')
    return f'=IF(AI{r}="","",IF({ld}=0,{base},{pick}))'


for k in range(N_SERIES * GRID):
    r = E_R0 + k
    sidx = f"AG{r}"
    m = f"AL{r}"
    ch_set = lambda c: f'AND(AM{r}="{CHANGE}",{chcol(c, m)}<>"")'  # noqa: E731
    f = {
        "A": "R",
        "AG": f"={k // GRID + 1}",
        "AH": f"=GridStart+{k % GRID}",
        "AI": f'=IF(OR({rcol("B", sidx)}="",{rcol("I", sidx)}=""),"",{rcol("A", sidx)})',
        "AJ": f'=IF(AI{r}="",FALSE,{occurs_expr(sidx, f"AH{r}")})',
        "AK": f'=IF(AJ{r},AI{r}&"|"&TEXT(AH{r},"0"),"")',
        "AL": f'=IF(AJ{r},IFERROR(MATCH(AK{r},Changes!$L${FIRST}:$L${C_LAST},0),0),0)',
        "AM": f'=IF(AL{r}=0,"",{chcol("C", m)})',
        "AN": from_override(r, "D", f'{rcol("C", sidx)}&""'),
        "AO": from_override(r, "E", rcol("D", sidx)),
        "AP": from_override(r, "F", rcol("E", sidx)),
        "AQ": from_override(r, "G", f'IF({rcol("H", sidx)}="","Confirmed",{rcol("H", sidx)})'),
        "AR": from_override(r, "H", f'IF({rcol("G", sidx)}="","Regular",{rcol("G", sidx)})'),
        "AS": f'=IF({ch_set("E")},{chcol("E", m)},AO{r})',
        "AT": f'=IF({ch_set("F")},{chcol("F", m)},AP{r})',
        "B": f'=IF(N{r},AI{r}&" "&TEXT(AH{r},"yyyy-mm-dd"),"")',
        "C": f'=IF(AJ{r},AH{r},"")',
        "D": f'=IF(AJ{r},{rcol("B", sidx)}&"","")',
        "E": f'=IF(OR(NOT(AJ{r}),J{r}="Unfilled"),"",IF({ch_set("D")},{chcol("D", m)},AN{r})&"")',
        "H": f'=AND(AJ{r},{rcol("F", sidx)}<>"Y",AS{r}<>"",AT{r}<>"",AS{r}<>AT{r})',
        "F": f'=IF(H{r},AS{r},"")',
        "G": f'=IF(H{r},AT{r},"")',
        "I": (f'=IF({ch_set("H")},{chcol("H", m)},IF(AND({ch_set("D")},{chcol("D", m)}<>AN{r}),"Coverage",AR{r}))'),
        "J": (f'=IF(AM{r}="{RELEASE}","Released",IF({ch_set("G")},{chcol("G", m)},AQ{r}))'),
        "K": (f'=IF(AND({ch_set("D")},{chcol("D", m)}<>AN{r},AN{r}<>""),AN{r},IF(AJ{r},{rcol("T", sidx)}&"",""))'),
        "L": "",
        "M": f'=IF(AND(AM{r}<>"",{chcol("I", m)}<>""),{chcol("I", m)}&"",IF(AJ{r},{rcol("U", sidx)}&"",""))',
        "N": f'=AND(AJ{r},AM{r}<>"{CANCEL}")',
    }
    f.update(common(r))
    for c, v in f.items():
        if c in ("AN", "AO", "AP", "AQ", "AR"):
            en[f"{c}{r}"] = ArrayFormula(f"{c}{r}", v)
        else:
            en[f"{c}{r}"] = v
for c in ("C", "AH"):
    for r in range(E_S0, E_LAST + 1):
        en[f"{c}{r}"].number_format = "yyyy-mm-dd"
for c in ("F", "G", "AO", "AP", "AS", "AT"):
    for r in range(E_S0, E_LAST + 1):
        en[f"{c}{r}"].number_format = "h:mm"
for c in ("O", "P", "X"):
    for r in range(E_S0, E_LAST + 1):
        en[f"{c}{r}"].number_format = "yyyy-mm-dd h:mm"
en.sheet_state = "hidden"

# ---------------------------------------------------------------------------
# _Slots: staffed count every 30 minutes, per requirement, only inside windows
# ---------------------------------------------------------------------------
sl = ws_sl
sl["A1"], sl["B1"], sl["C1"], sl["D1"] = "k", "t", "day", "slot"
for q in range(N_REQ):
    sl.cell(1, 5 + 2 * q, f"Req{q + 1} conf")
    sl.cell(1, 6 + 2 * q, f"Req{q + 1} any")
for k in range(N_SLOTS):
    r = 2 + k
    sl.cell(r, 1, k)
    sl.cell(r, 2, f"=GridStart+A{r}/48")
    sl.cell(r, 3, f"=INT(A{r}/48)")
    sl.cell(r, 4, f"=MOD(A{r},48)")
    for q in range(N_REQ):
        sr = REQ_FIRST + q
        applies = lambda day: (f'AND(INDEX(Setup!$K${sr}:$Q${sr},WEEKDAY(GridStart+{day}))="Y",'  # noqa: E731
                               f'GridStart+{day}>=N(Setup!$Z${sr}))')
        inside = (f'AND(Setup!$I${sr}<>"",Setup!$U${sr}<>"No",Setup!$X${sr}<>"",'
                  f'OR(AND({applies(f"C{r}")},D{r}>=Setup!$V${sr},OR(D{r}<Setup!$W${sr},Setup!$X${sr})),'
                  f'AND(Setup!$X${sr},{applies(f"C{r}-1")},D{r}<Setup!$W${sr})))')
        base = f'{E("D")},Setup!$I${sr},{E("O")},"<="&(B{r}+{EPS}),{E("P")},">"&(B{r}+{EPS})'
        sl.cell(r, 5 + 2 * q, f'=IF({inside},COUNTIFS({base},{E("R")},TRUE),"")')
        sl.cell(r, 6 + 2 * q, f'=IF({inside},COUNTIFS({base},{E("S")},TRUE),"")')
sl.sheet_state = "hidden"
SLOT_BLOCK = f"_Slots!$E$2:${L(4 + 2 * N_REQ)}${N_SLOTS + 1}"

# ---------------------------------------------------------------------------
# _Req: one row per (day in window, requirement)
# ---------------------------------------------------------------------------
rq = ws_rq
RQ_HEAD = ["day", "req", "date", "setup row", "applies", "slot a", "slot b", "min conf", "min any", "met",
           "met w/ tentative", "first short", "facility", "fac ok", "short", "key", "rank", "text"]
for i, h in enumerate(RQ_HEAD):
    rq.cell(1, 1 + i, h)
N_RQ_ROWS = GRID * N_REQ
for k in range(N_RQ_ROWS):
    r = 2 + k
    j, q = divmod(k, N_REQ)
    sr = REQ_FIRST + q
    rq.cell(r, 1, j)
    rq.cell(r, 2, q + 1)
    rq.cell(r, 3, f"=GridStart+A{r}")
    rq.cell(r, 3).number_format = "yyyy-mm-dd"
    rq.cell(r, 4, sr)
    rq.cell(r, 5, f'=AND(Setup!$I${sr}<>"",Setup!$U${sr}<>"No",Setup!$X${sr}<>"",N(Setup!$T${sr})>0,'
                  f'INDEX(Setup!$K${sr}:$Q${sr},WEEKDAY(C{r}))="Y",C{r}>=N(Setup!$Z${sr}))')
    rq.cell(r, 6, f'=IF(E{r},A{r}*48+Setup!$V${sr},"")')
    rq.cell(r, 7, f'=IF(E{r},F{r}+Setup!$Y${sr}-1,"")')
    col_c, col_a = 2 * q + 1, 2 * q + 2
    rq.cell(r, 8, f'=IF(E{r},MIN(INDEX({SLOT_BLOCK},F{r}+1,{col_c}):INDEX({SLOT_BLOCK},G{r}+1,{col_c})),"")')
    rq.cell(r, 9, f'=IF(E{r},MIN(INDEX({SLOT_BLOCK},F{r}+1,{col_a}):INDEX({SLOT_BLOCK},G{r}+1,{col_a})),"")')
    rq.cell(r, 10, f'=IF(E{r},H{r}>=Setup!$T${sr},"")')
    rq.cell(r, 11, f'=IF(E{r},I{r}>=Setup!$T${sr},"")')
    rq[f"L{r}"] = ArrayFormula(
        f"L{r}", f'=IF(AND(E{r},J{r}=FALSE),GridStart+(F{r}+MATCH(TRUE,INDEX({SLOT_BLOCK},F{r}+1,{col_c}):'
                 f'INDEX({SLOT_BLOCK},G{r}+1,{col_c})<Setup!$T${sr},0)-1)/48,"")')
    rq.cell(r, 12).number_format = "yyyy-mm-dd h:mm"
    rq.cell(r, 13, f'=Setup!$I${sr}&""')
    rq.cell(r, 14, f'=OR(FacFilter="All",FacFilter="",M{r}=FacFilter)')
    rq.cell(r, 15, f'=AND(E{r},J{r}=FALSE,N{r})')
    rq.cell(r, 16, f"=C{r}+B{r}/100")
    rq.cell(r, 17, f'=IF(O{r},COUNTIFS($O$2:$O${N_RQ_ROWS + 1},TRUE,$P$2:$P${N_RQ_ROWS + 1},"<"&P{r})+1,"")')
    rq.cell(r, 18, f'=IF(NOT(O{r}),"",M{r}&IF(Setup!$J${sr}<>""," ("&Setup!$J${sr}&")","")&": needs "&Setup!$T${sr}'
                   f'&", has "&H{r}&" confirmed from "&SUBSTITUTE(SUBSTITUTE(TEXT(L{r},"h:mm AM/PM")," AM","a")," PM","p")'
                   f'&IF(K{r}=TRUE," — met if tentative shifts are confirmed",""))')
rq.sheet_state = "hidden"
RQ = lambda c: f"_Req!${c}$2:${c}${N_RQ_ROWS + 1}"  # noqa: E731

# ---------------------------------------------------------------------------
# Agenda (sorted list for the calendar window, filtered)
# ---------------------------------------------------------------------------
ag = ws_ag
title(ag, "Agenda", "Every assignment in the six weeks shown on the Calendar, in time order, using the Calendar's filters.")
ag["A3"] = '="Showing "&TEXT(GridStart,"mmm d")&" – "&TEXT(GridEnd,"mmm d, yyyy")&IF(FacFilter="All",""," · "&FacFilter)&IF(NurseFilter="All",""," · "&NurseFilter)'
ag["A3"].font = font(size=10, bold=True)
AG_COLS = ["#", "Date", "Hours", "Nurse", "Facility", "Type", "Status", "Covering for", "Notes", "Coverage / warnings",
           "ID", "Calendar text"]
header(ag, 5, AG_COLS, [5, 15, 20, 20, 18, 11, 11, 16, 24, 40, 18, 10])
style_rows(ag, 6, 5 + N_AGENDA, 1, 12, wrap=True)
for k in range(N_AGENDA):
    r = 6 + k
    idx = f"$A{r}"
    ag.cell(r, 1, f'=IFERROR(MATCH({k + 1},{E("AC")},0),"")')
    ag.cell(r, 1).font = font(size=8, color="FFFFFF")
    g = lambda c: f'IF({idx}="","",INDEX({E(c)},{idx}))'  # noqa: E731
    ag.cell(r, 2, f"={g('C')}")
    ag.cell(r, 2).number_format = "ddd mmm d"
    tf = lambda x: f'TEXT({x},"h:mm AM/PM")'  # noqa: E731
    ag.cell(r, 3, f'=IF({idx}="","",IF(INDEX({E("H")},{idx}),{tf(f"INDEX({E(chr(70))},{idx})")}&"–"&'
                  f'{tf(f"INDEX({E(chr(71))},{idx})")}&IF(INDEX({E("G")},{idx})<=INDEX({E("F")},{idx})," (next day)",""),"Not specified"))')
    ag.cell(r, 4, f'=IF({idx}="","",IF(INDEX({E("J")},{idx})="Unfilled","UNFILLED",INDEX({E("E")},{idx})))')
    ag.cell(r, 5, f"={g('D')}")
    ag.cell(r, 6, f"={g('I')}")
    ag.cell(r, 7, f"={g('J')}")
    ag.cell(r, 8, f"={g('K')}")
    ag.cell(r, 9, f"={g('M')}")
    ag.cell(r, 10, f'=IF({idx}="","",_xlfn.TEXTJOIN("; ",TRUE,INDEX({E("AE")},{idx}),INDEX({E("T")},{idx})))')
    ag.cell(r, 11, f"={g('B')}")
    ag.cell(r, 12, f"={g('Y')}")
ag.column_dimensions["L"].hidden = True
ag.freeze_panes = "B6"
red = PatternFill("solid", fgColor="FDE2E2")
amber = PatternFill("solid", fgColor="FFF1D6")
violet = PatternFill("solid", fgColor="EFE7FD")
blue = PatternFill("solid", fgColor="E6EEFD")
grey = PatternFill("solid", fgColor="EEEFF1")
rng_ag = f"B6:K{5 + N_AGENDA}"
ag.conditional_formatting.add(rng_ag, FormulaRule(formula=['OR($G6="Unfilled",AND($G6="Released",LEFT($J6,5)="Still"))'], fill=red, font=Font(name=FONT, color="8F1414", bold=True), stopIfTrue=True))
ag.conditional_formatting.add(rng_ag, FormulaRule(formula=['$G6="Released"'], fill=grey, font=Font(name=FONT, color="6B7280", strike=True), stopIfTrue=True))
ag.conditional_formatting.add(rng_ag, FormulaRule(formula=['$G6="Tentative"'], fill=amber, font=Font(name=FONT, italic=True), stopIfTrue=True))
ag.conditional_formatting.add(rng_ag, FormulaRule(formula=['$F6="Coverage"'], fill=violet, stopIfTrue=True))
ag.conditional_formatting.add(rng_ag, FormulaRule(formula=['$F6="Regular"'], fill=blue, stopIfTrue=True))

# ---------------------------------------------------------------------------
# Attention
# ---------------------------------------------------------------------------
at = ws_att
title(at, "Needs attention",
      "Shifts marked Unfilled or Released that still need coverage, and unmet staffing requirements — for the six weeks on the Calendar.")
at["A3"] = ag["A3"].value.replace('IF(NurseFilter="All",""," · "&NurseFilter)', '""')
at["A3"].font = font(size=10, bold=True)
at["A4"], at["H4"] = "Unfilled shifts", "Below staffing requirement"
at["A4"].font = at["H4"].font = font(size=12, bold=True, color="8F1414")
header(at, 5, ["Date", "Hours", "Facility", "Original nurse", "Coverage still needed", "ID (for 'Covers shift ID')"],
       [14, 22, 18, 18, 46, 18])
header(at, 5, ["Date", "Shortfall"], [14, 70], start_col=8)
style_rows(at, 6, 5 + N_ATT, 1, 6, wrap=True)
style_rows(at, 6, 5 + N_ATT, 8, 9, wrap=True)
for k in range(N_ATT):
    r = 6 + k
    idx = f'IFERROR(MATCH({k + 1},{E("AF")},0),"")'
    at.cell(r, 15, f"={idx}")  # helper column O
    i = f"$O{r}"
    g = lambda c: f'IF({i}="","",INDEX({E(c)},{i}))'  # noqa: E731
    at.cell(r, 1, f"={g('C')}")
    at.cell(r, 1).number_format = "ddd mmm d"
    at.cell(r, 2, f'=IF({i}="","",IF(INDEX({E("H")},{i}),TEXT(INDEX({E("F")},{i}),"h:mm AM/PM")&"–"&'
                  f'TEXT(INDEX({E("G")},{i}),"h:mm AM/PM"),"Not specified"))')
    at.cell(r, 3, f"={g('D')}")
    at.cell(r, 4, f'=IF({i}="","",IF(INDEX({E("J")},{i})="Unfilled","(unfilled shift)",INDEX({E("E")},{i})&" — released"))')
    at.cell(r, 5, f"={g('AE')}")
    at.cell(r, 6, f"={g('B')}")
    ridx = f'IFERROR(MATCH({k + 1},{RQ("Q")},0),"")'
    at.cell(r, 16, f"={ridx}")  # helper column P
    at.cell(r, 8, f'=IF($P{r}="","",INDEX({RQ("C")},$P{r}))')
    at.cell(r, 8).number_format = "ddd mmm d"
    at.cell(r, 9, f'=IF($P{r}="","",INDEX({RQ("R")},$P{r}))')
at.column_dimensions["O"].hidden = True
at.column_dimensions["P"].hidden = True
at["A" + str(7 + N_ATT)] = "Blank days are not gaps: only shifts marked Unfilled/Released and configured requirements are flagged."
at["A" + str(7 + N_ATT)].font = font(size=9, italic=True, color=MUTED)

# ---------------------------------------------------------------------------
# Calendar
# ---------------------------------------------------------------------------
title(cal, "Nurse Schedule", "Type any date in the month box to change months. Add shifts on the Shifts and Recurring sheets.")
cal["B3"], cal["D3"], cal["F3"] = "Month:", "Facility:", "Nurse:"
for c in ("B3", "D3", "F3"):
    cal[c].font = font(bold=True)
    cal[c].alignment = Alignment(horizontal="right")
cal["C3"] = "=TODAY()"
cal["C3"].number_format = "mmmm yyyy"
cal["E3"], cal["G3"] = "All", "All"
for c in ("C3", "E3", "G3"):
    cal[c].fill = INPUT_FILL
    cal[c].font = font(bold=True, color=BLUE_IN, size=12)
    cal[c].border = BOX
cal["C3"].comment = Comment("Type any date (e.g. 12/1/2026) to show that month. =TODAY() shows the current month.", "Nurse Schedule")
add_list(cal, "E3", f"=_Lists!$M$2:$M${2 + N_FAC}")
add_list(cal, "G3", f"=_Lists!$N$2:$N${2 + N_NURSES}")
cal["L1"] = "=DATE(YEAR(MonthPick),MONTH(MonthPick),1)-WEEKDAY(DATE(YEAR(MonthPick),MONTH(MonthPick),1))+1"
cal["M1"] = f"=L1+{GRID - 1}"
cal["B5"] = '=TEXT(MonthPick,"mmmm yyyy")'
cal["B5"].font = font(size=16, bold=True)
cal["D5"] = '="Showing "&TEXT(GridStart,"mmm d")&" – "&TEXT(GridEnd,"mmm d, yyyy")&IF(E3="All",""," · "&E3)&IF(G3="All",""," · "&G3)'
cal["D5"].font = font(size=9, color=MUTED, italic=True)
for i, d in enumerate(["Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"]):
    c = cal.cell(6, 2 + i, d)
    c.font = font(bold=True, size=10, color=MUTED)
    c.alignment = Alignment(horizontal="center")
    c.fill = HEAD_FILL
    cal.column_dimensions[L(2 + i)].width = 30
cal.column_dimensions["A"].width = 2
AGL = f"Agenda!$L$6:$L${5 + N_AGENDA}"
AGD = f"Agenda!$B$6:$B${5 + N_AGENDA}"
for w in range(6):
    dr, er = 7 + 2 * w, 8 + 2 * w
    cal.row_dimensions[er].height = 118
    for dcol in range(7):
        col = L(2 + dcol)
        hcol = L(12 + dcol)   # hidden helpers L..R
        off = w * 7 + dcol
        dc = cal[f"{col}{dr}"]
        dc.value = f"=GridStart+{off}"
        dc.number_format = "d"
        dc.font = font(bold=True, size=11)
        dc.alignment = Alignment(horizontal="left")
        dc.border = Border(left=THIN, right=THIN, top=THIN)
        # helper: requirement badge + open flag
        cal[f"{hcol}{dr}"] = (f'=IF(COUNTIFS({RQ("C")},{col}{dr},{RQ("E")},TRUE,{RQ("N")},TRUE)=0,"",'
                              f'IF(COUNTIFS({RQ("C")},{col}{dr},{RQ("O")},TRUE)=0,"✓ Staffed",'
                              f'IF(COUNTIFS({RQ("C")},{col}{dr},{RQ("O")},TRUE,{RQ("K")},FALSE)=0,"◐ Short unless tentative confirmed","✕ Short")))')
        cal[f"{hcol}{er}"] = f'=COUNTIFS({E("C")},{col}{dr},{E("AD")},TRUE,{E("AA")},TRUE)>0'
        ec = cal[f"{col}{er}"]
        ec.value = ArrayFormula(f"{col}{er}", f'=_xlfn.TEXTJOIN(CHAR(10),TRUE,{hcol}{dr},IF({AGD}={col}{dr},{AGL},""))')
        ec.alignment = Alignment(wrap_text=True, vertical="top")
        ec.font = font(size=9)
        ec.border = Border(left=THIN, right=THIN, bottom=THIN)
    for h in range(12, 19):
        cal.column_dimensions[L(h)].hidden = True
grid_dates = "B7:H7 B9:H9 B11:H11 B13:H13 B15:H15 B17:H17"
grid_text = "B8:H8 B10:H10 B12:H12 B14:H14 B16:H16 B18:H18"
cal.conditional_formatting.add(grid_dates, FormulaRule(formula=["B7=TODAY()"], fill=PatternFill("solid", fgColor="FFF3B0")))
cal.conditional_formatting.add(grid_dates, FormulaRule(formula=["MONTH(B7)<>MONTH(MonthPick)"], font=Font(name=FONT, color="A0A7B2")))
cal.conditional_formatting.add(grid_text, FormulaRule(formula=["L8=TRUE"], fill=PatternFill("solid", fgColor="FDE2E2"), stopIfTrue=True))
cal.conditional_formatting.add(grid_text, FormulaRule(formula=['LEFT(L7,1)="✕"'], fill=PatternFill("solid", fgColor="FFF0F0"), stopIfTrue=True))
cal.conditional_formatting.add(grid_text, FormulaRule(formula=["MONTH(B7)<>MONTH(MonthPick)"], fill=PatternFill("solid", fgColor="F7F7F8")))
cal["B20"] = ("Legend: each line is “hours  nurse · facility · type”. REGULAR / Coverage / UNFILLED / RELEASED labels, "
              "“Tentative” for unconfirmed, “NEEDS COVERAGE” and red day = open time, “⚠ overlap” = nurse double-booked. "
              "Top line: ✓ Staffed / ✕ Short / ◐ Short unless tentative confirmed (only on days with a requirement).")
cal["B20"].font = font(size=9, italic=True, color=MUTED)
cal.merge_cells("B20:H21")
cal["B20"].alignment = Alignment(wrap_text=True, vertical="top")
cal.freeze_panes = "A7"
cal.sheet_view.showGridLines = False

# ---------------------------------------------------------------------------
# Instructions
# ---------------------------------------------------------------------------
hp = ws_help
hp.column_dimensions["A"].width = 3
hp.column_dimensions["B"].width = 110
lines = [
    ("Nurse Schedule — how to use", "h1"),
    ("Yellow cells with blue text are where you type. Grey cells are calculated. Works in Excel for Microsoft 365 (desktop or web); no macros.", "p"),
    ("Don't sort or delete rows on Shifts, Recurring or Changes — clear a row's contents instead, so IDs stay put.", "warn"),
    ("1. Set up", "h2"),
    ("Setup sheet: list your facilities and nurses. Set a nurse's Active to No to hide them from new assignments (their history stays).", "p"),
    ("Optional: add staffing requirements (facility, weekdays, time window, minimum nurses, and an optional date they start applying). Without them, empty days are never flagged.", "p"),
    ("2. Add a one-time shift or coverage (Shifts sheet)", "h2"),
    ("Fill Date, Facility, Nurse, Start, End. Type: Regular (default) or Coverage. Status: Confirmed (default), Tentative, Unfilled.", "p"),
    ("Overnight: an End earlier than Start means it ends the next day (shown in 'Ends next day'). If times aren't known, put Y in 'Hours not specified' — it is never treated as 24-hour coverage.", "p"),
    ("Same assignment on several dates: copy the row and change the date (Ctrl+D fills down). Duplicate to another date the same way.", "p"),
    ("Unfilled shift: set Status to Unfilled and leave Nurse blank.", "p"),
    ("3. Recurring schedules (Recurring sheet)", "h2"),
    ("Enter facility, nurse, hours, first date, 'Every N weeks' (1 = weekly, 2 = every other week, 3 = every 3 weeks…), and Y under each weekday.", "p"),
    ("End: leave blank for no end date, set an End date, or 'Stop after N times'. Check the Summary column, e.g. “Every other week on Monday, Wednesday, 8:00 AM–4:00 PM, starting November 2, 2026, through December 31, 2026.”", "p"),
    ("4. Change a recurring schedule (Changes sheet)", "h2"),
    ("This occurrence only: Schedule ID + Date + 'Change this date', then fill only the fields that change (e.g. New nurse). The rest of the series is untouched. Giving a different nurse marks it Coverage, covering for the original nurse.", "p"),
    ("Cancel one date: 'Cancel this date'.", "p"),
    ("This and future occurrences: 'Change from this date on' with the new values. Earlier dates keep their old values; one-off changes stay in place.", "p"),
    ("Entire series: edit the row on the Recurring sheet directly. One-off changes on the Changes sheet are kept.", "p"),
    ("Stop a series from a date: set its End date on the Recurring sheet.", "p"),
    ("The Check column confirms the date is really part of that schedule.", "p"),
    ("5. Coverage", "h2"),
    ("Full replacement on a recurring date: Changes → 'Change this date' → New nurse. On a one-time shift: just change the nurse, or set the original row's Status to Released and add a coverage row.", "p"),
    ("Partial coverage: mark the original Released (one-time: Status = Released; recurring: Changes → 'Release this date'). Then add Coverage rows on the Shifts sheet with the hours covered and the original's ID in 'Covers shift ID' (S12, or the occurrence ID like 'R3 2026-11-30' from the Changes or Attention sheet). The released nurse is no longer counted and is filled in as 'Covering for'. The remaining unfilled hours appear in 'Coverage still needed' and on the Attention sheet.", "p"),
    ("Create a coverage assignment without an existing shift: just add a Shifts row with Type = Coverage (Covering for optional).", "p"),
    ("6. Reading the schedule", "h2"),
    ("Calendar: six weeks around the month you pick; filter by facility and nurse. Each line shows hours, nurse, facility and type. Days with open coverage turn red. Staffing badges appear only on days with a requirement.", "p"),
    ("Agenda: the same assignments as a sorted, colour-coded list. Attention: unfilled time and unmet requirements.", "p"),
    ("Overlaps: the Warnings column (Shifts) and '⚠ overlap' (Calendar) flag the same nurse booked twice, across facilities and overnight. Saving is never blocked — an intentional double booking just stays flagged.", "p"),
    ("Limits", "h2"),
    (f"Up to {N_SHIFTS} one-time shifts, {N_SERIES} recurring schedules, {N_CHANGES} changes, {N_NURSES} nurses, {N_FAC} facilities, {N_REQ} requirements. "
     "Recurring dates are calculated for the six weeks shown on the Calendar, so overlap and coverage checks involving recurring shifts cover that window. "
     "All times are local wall-clock times (one timezone); requirements are checked every 30 minutes; a shift can't be moved to another date via Changes (cancel it and add a one-time shift).", "p"),
    ("Sharing: save the file to OneDrive or SharePoint and share it; several people can edit it in Excel for the web at the same time.", "p"),
]
r = 2
for text, kind in lines:
    c = hp.cell(r, 2, text)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if kind == "h1":
        c.font = font(size=16, bold=True, color="2563EB")
    elif kind == "h2":
        c.font = font(size=12, bold=True)
        r += 0
    elif kind == "warn":
        c.font = font(size=10, color="B91C1C", bold=True)
    else:
        c.font = font(size=10)
    r += 1
hp.sheet_view.showGridLines = False

# ---------------------------------------------------------------------------
# Example rows (delete or overwrite them)
# ---------------------------------------------------------------------------
if "--empty" not in sys.argv:
    today = dt.date.today()
    mon = today - dt.timedelta(days=today.weekday())  # Monday this week
    st["A6"], st["B6"] = "Ana Ruiz (example)", "Y"
    st["A7"], st["B7"] = "Ben Okafor (example)", "Y"
    st["A8"], st["B8"] = "Cara Lin (example)", "Y"
    st["F6"] = "North Clinic (example)"
    st["F7"] = "West Campus (example)"
    for c, v in zip("IJKLMNOPQRSTU", ["North Clinic (example)", "Day shift", "", "Y", "Y", "Y", "Y", "Y", "",
                                       dt.time(8, 0), dt.time(16, 0), 1, "Y"]):
        if v != "":
            st[f"{c}6"] = v
    st["Z6"] = mon
    rc["B6"], rc["C6"], rc["D6"], rc["E6"] = "North Clinic (example)", "Ana Ruiz (example)", dt.time(8), dt.time(16)
    rc["I6"], rc["J6"], rc["L6"], rc["N6"] = mon, 2, "Y", "Y"
    rc["U6"] = "Example — every other Mon & Wed"
    sh["B6"], sh["C6"], sh["D6"], sh["E6"], sh["F6"] = mon + dt.timedelta(days=1), "North Clinic (example)", "Ben Okafor (example)", dt.time(8), dt.time(16)
    sh["H6"], sh["I6"], sh["L6"] = "Coverage", "Cara Lin (example)", "Example coverage shift"
    sh["B7"], sh["C7"], sh["E7"], sh["F7"], sh["J7"] = mon + dt.timedelta(days=3), "North Clinic (example)", dt.time(8), dt.time(16), "Unfilled"
    sh["L7"] = "Example unfilled shift"
    cg["A6"], cg["B6"], cg["C6"], cg["D6"] = "R1", mon + dt.timedelta(days=14), CHANGE, "Ben Okafor (example)"
    cg["I6"] = "Example — Ben covers this one date"

# Excel computes everything when the file is opened.
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("wrote", OUT)
